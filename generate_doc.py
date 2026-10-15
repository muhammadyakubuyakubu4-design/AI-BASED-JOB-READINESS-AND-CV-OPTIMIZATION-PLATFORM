import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('CareerAI (JOB) - Project Defense Guide', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('A comprehensive overview of the system architecture, features, and underlying logic for your project defense.')

# Section 1
doc.add_heading('1. System Overview (The Big Picture)', level=1)
doc.add_paragraph('CareerAI is a comprehensive, AI-powered career development platform that acts as a personal career coach. The system is divided into several core modules:')

doc.add_paragraph('CV Analyzer & Editor: Uploads resumes, extracts text, calculates performance scores, provides AI-driven suggestions, and dynamically edits the user\'s MS Word document in real-time.', style='List Bullet')
doc.add_paragraph('Job Matching Engine: Compares the user\'s CV against a specific Job Description to calculate a targeted match percentage.', style='List Bullet')
doc.add_paragraph('AI Career Advisor: Analyzes the user\'s skills to recommend career paths, generate study plans, suggest job titles, and link directly to real-world hiring companies.', style='List Bullet')
doc.add_paragraph('Learning & Assessment Hub: Provides AI-generated lessons, quizzes, and study plans to help users upskill.', style='List Bullet')

# Section 2
doc.add_heading('2. How the CV Analysis Works (Step-by-Step)', level=1)
doc.add_paragraph('When a user uploads a CV, the system follows a complex pipeline to analyze it:')

doc.add_paragraph('Step 1 - Document Parsing: The system takes the PDF or DOCX file and uses backend libraries (pdfplumber for PDFs, python-docx for Word docs) to extract the raw text.', style='List Number')
doc.add_paragraph('Step 2 - NLP Section Extraction: The system uses regular expressions (Regex) and Natural Language Processing (NLP) to break the document down into logical categories (Contact, Summary, Experience, Education, Skills).', style='List Number')
doc.add_paragraph('Step 3 - Algorithmic Scoring: The system passes the text through mathematical heuristic engines to calculate the ATS, Grammar, Format, and Keyword scores out of 100.', style='List Number')
doc.add_paragraph('Step 4 - AI Feedback Generation: The parsed text is securely sent to a Large Language Model (Groq / Qwen AI). The AI generates a qualitative summary and high-priority improvement suggestions.', style='List Number')
doc.add_paragraph('Step 5 - Dynamic DOCX Updating: If the user clicks \'Enhance\', the AI rewrites it. The backend Python server then physically opens the uploaded .docx file, locates the exact paragraph, replaces it with the newly enhanced text, saves the file, and forces the browser to refresh.', style='List Number')

# Section 3
doc.add_heading('3. How the Keyword Engine Works', level=1)
doc.add_paragraph('The Keyword Checker operates in two different modes depending on what the user is doing:')

doc.add_paragraph('General Mode (No Job Description): Scans the text against a hardcoded dictionary of high-value industry action verbs and nouns (e.g., team, project, develop, manage, implement, design, analyze, lead, communicate, deliver).', style='List Bullet')
doc.add_paragraph('Targeted Mode (Job Match Tab): When the user pastes a Job Description, the system strips out all meaningless stopwords (the, a, and, but). It then calculates the exact mathematical overlap/intersection of the remaining meaningful words to generate a targeted match percentage.', style='List Bullet')

# Section 4
doc.add_heading('4. How the Grammar & Professionalism Check Works', level=1)
doc.add_paragraph('Rather than just doing basic spell-check, the system runs a Heuristic Professionalism Algorithm:')

doc.add_paragraph('Formatting Anomalies: Scans for highly fragmented sentences or very short, broken lines that indicate poor formatting or copy-paste errors.', style='List Bullet')
doc.add_paragraph('Excessive Capitalization: Searches for long lines of text that are entirely uppercase and penalizes them to prevent the CV from shouting.', style='List Bullet')
doc.add_paragraph('Informal Vocabulary Check: Scans the text for colloquial or unprofessional slang words (e.g., gonna, wanna, kinda, sorta, stuff, things).', style='List Bullet')
doc.add_paragraph('Vocabulary Diversity: Calculates a Unique Word Ratio. If the user repeats the exact same words too frequently, the unique ratio drops below 40%, and the system deducts points for lacking descriptive variety.', style='List Bullet')

# Tip
doc.add_heading('Pro Tip for Your Defense:', level=2)
p = doc.add_paragraph('If the panel asks you why you didn\'t just use AI for everything, tell them: ')
p.add_run('\"Relying entirely on AI is slow and expensive. I designed a hybrid architecture. I use fast, mathematical Python algorithms for the quantitative scoring (Keywords, Grammar, ATS), and I reserve the heavy AI processing strictly for qualitative tasks like rewriting text and giving career advice. This makes the system incredibly fast, cost-effective, and scalable.\"').italic = True

output_path = r'C:\Users\HP\Desktop\CareerAI_Project_Defense.docx'
doc.save(output_path)
print(f'Document saved to: {output_path}')
