import os
import re
from django.conf import settings
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.views.decorators.clickjacking import xframe_options_exempt
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from .models import CVUpload, CVAnalysis
from .forms import CVUploadForm
from .utils.parser import parse_cv_file, extract_sections, extract_contact_info
from .utils.ai_service import analyze_cv, analyze_bullet_points



@login_required
def cv_upload_view(request):
    """Handle CV file upload and trigger analysis."""
    if request.method == 'POST':
        form = CVUploadForm(request.POST, request.FILES)
        if form.is_valid():
            cv_file = request.FILES['file']
            ext = cv_file.name.rsplit('.', 1)[-1].lower()

            # Save the uploaded file
            cv_upload = CVUpload(
                user=request.user,
                file=cv_file,
                original_filename=cv_file.name,
                file_type=ext,
                file_size=cv_file.size,
            )
            cv_upload.save()

            # Set all other CVs as inactive
            CVUpload.objects.filter(user=request.user).exclude(pk=cv_upload.pk).update(is_active=False)

            # Parse the CV text
            try:
                file_path = cv_upload.file.path
                parsed_text = parse_cv_file(file_path)
                cv_upload.parsed_text = parsed_text
                cv_upload.save()

                # Run AI analysis
                analysis_result = analyze_cv(parsed_text)

                # Run bullet point feedback (Feature 5)
                try:
                    bullet_fb = analyze_bullet_points(parsed_text)
                except Exception:
                    bullet_fb = []

                # Save analysis
                CVAnalysis.objects.create(
                    cv_upload=cv_upload,
                    ats_score=analysis_result['ats_score'],
                    grammar_score=analysis_result['grammar_score'],
                    keyword_score=analysis_result['keyword_score'],
                    format_score=analysis_result['format_score'],
                    overall_score=analysis_result['overall_score'],
                    extracted_sections=analysis_result['extracted_sections'],
                    extracted_skills=analysis_result['extracted_skills'],
                    detected_keywords=analysis_result.get('detected_keywords', analysis_result.get('extracted_skills', [])),
                    contact_info=analysis_result['contact_info'],
                    suggestions=analysis_result['suggestions'],
                    missing_sections=analysis_result['missing_sections'],
                    improvements=analysis_result['improvements'],
                    ai_feedback=analysis_result['ai_feedback'],
                    bullet_feedback=bullet_fb,
                )

                # Log activity
                try:
                    from apps.admin_panel.utils import log_activity
                    log_activity(request.user, 'upload', f'Uploaded CV "{cv_upload.original_filename}" (score: {analysis_result["overall_score"]}/100)', 'CVUpload', cv_upload.pk, request)
                except Exception:
                    pass

                # Update user readiness score
                if hasattr(request.user, 'profile'):
                    request.user.profile.readiness_score = analysis_result['overall_score']
                    request.user.profile.save()

                # Create notification
                _create_cv_notification(request.user, cv_upload, analysis_result['overall_score'])

                # Auto-generate career/job recommendations based on CV skills
                try:
                    from apps.cv_analyzer.utils.ai_service import get_career_recommendations
                    from apps.career.models import CareerRecommendation
                    
                    target_role = getattr(request.user.profile, 'target_role', '') if hasattr(request.user, 'profile') else ''
                    experience = getattr(request.user.profile, 'years_experience', '0') if hasattr(request.user, 'profile') else '0'
                    
                    skills_list = analysis_result.get('extracted_skills', [])
                    skills = [{'skill_name': s} for s in set(skills_list)]
                    
                    if skills:
                        CareerRecommendation.objects.filter(user=request.user).update(is_current=False)
                        rec_res = get_career_recommendations(skills, target_role, experience)
                        CareerRecommendation.objects.create(
                            user=request.user,
                            recommended_skills=rec_res.get('recommended_skills', []),
                            recommended_courses=rec_res.get('recommended_courses', []),
                            career_paths=rec_res.get('career_paths', []),
                            certifications=rec_res.get('certifications', []),
                            job_titles=rec_res.get('job_titles', []),
                            target_companies=rec_res.get('target_companies', []),
                            learning_resources=rec_res.get('learning_resources', []),
                            current_skills_snapshot=[s['skill_name'] for s in skills],
                            target_role=target_role,
                            ai_summary=rec_res.get('ai_summary', ''),
                            is_current=True
                        )
                except Exception:
                    pass

                messages.success(request, f'CV uploaded and analyzed successfully! Overall score: {analysis_result["overall_score"]}/100')
                return redirect('cv_analyzer:analysis', pk=cv_upload.pk)

            except Exception as e:
                cv_upload.delete()
                messages.error(request, f'Failed to parse CV: {str(e)}. Please ensure the file is not corrupted.')
        else:
            messages.error(request, 'Please correct the form errors.')
    else:
        form = CVUploadForm()

    recent_cvs = CVUpload.objects.filter(user=request.user, is_deleted=False).order_by('-uploaded_at')[:5]
    context = {
        'page_title': 'Upload CV',
        'form': form,
        'recent_cvs': recent_cvs,
    }
    return render(request, 'cv/upload.html', context)


@login_required
def cv_analysis_view(request, pk):
    """Display CV analysis results with full tabbed interface."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)

    if not cv_upload.has_analysis:
        messages.warning(request, 'This CV has not been analyzed yet.')
        return redirect('cv_analyzer:upload')

    analysis = cv_upload.analysis

    from apps.assessment.models import UserSkill
    from apps.career.models import CareerRecommendation
    from apps.interview_prep.models import InterviewSession
    from apps.dashboard.models import UserProfile
    from django.db.models import Avg

    profile = getattr(request.user, 'profile', None)
    skills = UserSkill.objects.filter(user=request.user)
    recommendations = CareerRecommendation.objects.filter(user=request.user, is_current=True).first()

    avg_scores = CVAnalysis.objects.filter(cv_upload__user=request.user).aggregate(
        avg_overall=Avg('overall_score'), avg_ats=Avg('ats_score'),
        avg_grammar=Avg('grammar_score'), avg_keyword=Avg('keyword_score'),
        avg_format=Avg('format_score'),
    )

    context = {
        'page_title': 'CV Analysis',
        'cv_upload': cv_upload,
        'analysis': analysis,
        'parsed_text': cv_upload.parsed_text or '',
        'profile': profile,
        'skills': skills,
        'recommendations': recommendations,
        'avg_scores': {k: round(v or 0, 1) for k, v in avg_scores.items()},
        'score_labels': {
            'ats': _get_score_label(analysis.ats_score),
            'grammar': _get_score_label(analysis.grammar_score),
            'keyword': _get_score_label(analysis.keyword_score),
            'format': _get_score_label(analysis.format_score),
            'overall': _get_score_label(analysis.overall_score),
        },
        'bullet_feedback': analysis.bullet_feedback or [],
        'ai_provider': 'Groq' if getattr(settings, 'OPENAI_API_KEY', '').startswith('gsk_') else 'OpenAI',
    }
    return render(request, 'cv/analysis.html', context)



@login_required
def cv_enhance_view(request, pk):
    """AJAX: Enhance CV content with AI rewriting."""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    import json
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    data = json.loads(request.body)
    section = data.get('section', '')
    text = data.get('text', '')
    action = data.get('action', 'improve')

    if not text.strip():
        return JsonResponse({'error': 'No text provided'}, status=400)

    prompts = {
        'improve': 'Improve the following CV text to be more professional and impactful. Return ONLY the improved text without explanations.',
        'grammar': 'Fix all grammar, spelling, and punctuation issues in the following CV text. Return ONLY the corrected text.',
        'tone': 'Rewrite the following CV text with a more professional and confident tone. Return ONLY the rewritten text.',
        'clarity': 'Improve the clarity and readability of the following CV text. Make it concise and easy to understand. Return ONLY the improved text.',
        'impact': 'Make the following CV text more impactful by using stronger action verbs and quantifiable achievements. Return ONLY the improved text.',
        'bullet': 'Rewrite the following as powerful, achievement-oriented bullet points for a CV. Return ONLY the bullet points, one per line.',
        'summary': 'Write a compelling professional summary (2-3 sentences) for this CV. Use specific details from the person\'s experience, skills, and background — their actual job titles, industries, key skills, and notable achievements. Avoid generic phrases like "results-driven", "team player", "forward-thinking", "proven track record", or "passionate about". Be specific and factual. Return ONLY the summary.',
    }

    prompt = prompts.get(action, prompts['improve'])

    from .utils.ai_service import _get_openai_client, _get_gemini_model
    client = _get_openai_client()
    gemini = _get_gemini_model()

    # For summary action, use the CV parsed text as source instead of a command string
    source_text = text
    if action == 'summary':
        cv_text = cv_upload.parsed_text or ''
        # If the user text is just a generic command, replace with actual CV content
        if len(text.strip().split()) < 15:
            source_text = cv_text[:2500] or text
        else:
            source_text = text

    if client or gemini:
        sections_text = cv_upload.parsed_text[:1200] if cv_upload.parsed_text else ''
        combined = source_text if action == 'summary' else f'{source_text}\n\nContext from CV: {sections_text}'
        
        try:
            if gemini:
                full_prompt = f"You are a professional CV writer and career coach.\n\n{prompt}\n\nCRITICAL INSTRUCTION: Output ONLY the requested specific text. DO NOT rewrite the entire CV. DO NOT include markdown asterisks (**).\n\nText:\n{combined[:2500]}"
                response = gemini.generate_content(full_prompt)
                enhanced = response.text.strip()
            else:
                response = client.chat.completions.create(
                    model='qwen/qwen3.8-27b' if getattr(settings, 'OPENAI_API_KEY', '').startswith('gsk_') else 'gpt-4o-mini',
                    messages=[
                        {'role': 'system', 'content': 'You are a professional CV writer and career coach. CRITICAL INSTRUCTION: Output ONLY the requested specific text. DO NOT rewrite the entire CV. DO NOT include markdown asterisks (**).'},
                        {'role': 'user', 'content': f'{prompt}\n\nText:\n{combined[:2500]}'},
                    ],
                    max_tokens=1000,
                    temperature=0.7,
                )
                enhanced = response.choices[0].message.content.strip()
            
            import re
            enhanced = re.sub(r'^```[a-z]*\n?', '', enhanced)
            enhanced = re.sub(r'\n?```$', '', enhanced)

            # Update the original DOCX file if it's a docx
            if cv_upload.file_type == 'docx':
                try:
                    from docx import Document
                    import os
                    file_path = cv_upload.file.path
                    if os.path.exists(file_path):
                        doc = Document(file_path)
                        replaced = False
                        
                        # Clean up markdown syntax from AI response so it looks good in Word
                        clean_enhanced = enhanced.replace('**', '').replace('###', '').replace('##', '').replace('*', '')

                        if action == 'summary' and text.startswith('Generate'):
                            # The user is generating a brand new summary, insert it at the top of the CV
                            # We insert after the contact block (usually paragraph 1 or 2)
                            if len(doc.paragraphs) > 2:
                                p = doc.paragraphs[2].insert_paragraph_before()
                                p.add_run('PROFESSIONAL SUMMARY').bold = True
                                for line in clean_enhanced.split('\n'):
                                    if line.strip(): doc.paragraphs[2].insert_paragraph_before(line.strip())
                            else:
                                doc.add_paragraph('PROFESSIONAL SUMMARY')
                                for line in clean_enhanced.split('\n'):
                                    if line.strip(): doc.add_paragraph(line.strip())
                            replaced = True
                        else:
                            # Try to find the existing text to replace
                            lines = [line.strip() for line in text.split('\n') if line.strip() and len(line.strip()) > 10]
                            if lines:
                                for para in doc.paragraphs:
                                    if not para.text.strip(): continue
                                    if lines[0] in para.text or para.text in lines[0]:
                                        para.text = clean_enhanced
                                        replaced = True
                                        break
                                if not replaced:
                                    for table in doc.tables:
                                        for row in table.rows:
                                            for cell in row.cells:
                                                for para in cell.paragraphs:
                                                    if lines[0] in para.text or para.text in lines[0]:
                                                        para.text = clean_enhanced
                                                        replaced = True
                                                        break
                                                if replaced: break
                                        if replaced: break
                            
                            # If we still couldn't find the text to replace (meaning the user generated a completely new section),
                            # we append it to the end of the document automatically!
                            if not replaced:
                                doc.add_paragraph('')
                                p = doc.add_paragraph()
                                p.add_run('NEW ADDITION').bold = True
                                for line in clean_enhanced.split('\n'):
                                    if line.strip(): doc.add_paragraph(line.strip())
                                replaced = True

                        if replaced:
                            doc.save(file_path)
                except Exception as e:
                    pass # Ignore docx update errors

            # Update the parsed_text in the database so the UI is up to date
            if text.strip() and text in cv_upload.parsed_text:
                cv_upload.parsed_text = cv_upload.parsed_text.replace(text, enhanced)
                cv_upload.save(update_fields=['parsed_text'])

            return JsonResponse({'enhanced': enhanced})
        except Exception:
            pass

    # Offline deterministic enhancement fallback (no API required)
    from .utils.local_enhancer import enhance_local
    try:
        local_result = enhance_local(text if action != 'summary' else (text or ''), action)
        if local_result and local_result.strip():
            return JsonResponse({'enhanced': local_result})
    except Exception:
        pass
    return JsonResponse({'enhanced': text})


@login_required
def cv_match_jd_view(request, pk):
    """AJAX: Match CV against a job description."""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    import json
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    data = json.loads(request.body)
    job_text = data.get('job_description', '').strip()

    if not job_text:
        return JsonResponse({'error': 'Job description is required'}, status=400)

    cv_text = cv_upload.parsed_text or ''

    from .utils.analyzer import get_missing_keywords
    from .utils.parser import extract_skills_from_text

    keyword_analysis = get_missing_keywords(cv_text, job_text)
    match_score = keyword_analysis['match_percentage']

    cv_skills = extract_skills_from_text(cv_text)
    job_skills = extract_skills_from_text(job_text)
    missing_skills = [s for s in job_skills if s not in cv_skills]
    matching_skills = [s for s in job_skills if s in cv_skills]

    gap_percentage = round(len(missing_skills) / max(len(job_skills), 1) * 100, 1)

    suggestions = []
    if missing_skills:
        suggestions.append(f"Add these skills: {', '.join(missing_skills[:8])}")
    if keyword_analysis['missing']:
        suggestions.append(f"Include keywords: {', '.join(keyword_analysis['missing'][:8])}")
    if match_score < 40:
        suggestions.append('Consider tailoring your CV more closely to this job description.')
    elif match_score < 70:
        suggestions.append('Good foundation! Strengthen the areas mentioned above.')
    else:
        suggestions.append('Strong match! Minor tweaks to optimize further.')

    from .utils.ai_service import _get_openai_client
    client = _get_openai_client()
    ai_summary = ''
    if client:
        try:
            resp = client.chat.completions.create(
                model='qwen/qwen3.8-27b' if getattr(settings, 'OPENAI_API_KEY', '').startswith('gsk_') else 'gpt-4o-mini',
                messages=[
                    {'role': 'system', 'content': 'You are a recruitment expert analyzing CV-job fit.'},
                    {'role': 'user', 'content': f'CV text:\n{cv_text[:1500]}\n\nJob Description:\n{job_text[:1500]}\n\nProvide a brief assessment of fit (2-3 sentences) and 3 key recommendations.'},
                ],
                max_tokens=400,
                temperature=0.5,
            )
            ai_summary = resp.choices[0].message.content.strip()
        except Exception:
            pass

    return JsonResponse({
        'match_score': round(match_score, 1),
        'matching_keywords': keyword_analysis['matching'][:15],
        'missing_keywords': keyword_analysis['missing'][:15],
        'matching_skills': matching_skills,
        'missing_skills': missing_skills,
        'gap_percentage': gap_percentage,
        'suggestions': suggestions[:5],
        'ai_summary': ai_summary,
        'readiness_indicator': 'high' if match_score >= 70 else 'medium' if match_score >= 40 else 'low',
    })


@login_required
def cv_save_profile_view(request, pk):
    """AJAX: Save user profile edits from the CV analysis page."""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    import json
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    data = json.loads(request.body)
    section = data.get('section', '')
    value = data.get('value', '')

    profile = getattr(request.user, 'profile', None)
    if not profile:
        return JsonResponse({'error': 'Profile not found'}, status=404)

    field_map = {
        'personal': {'bio': value},
        'objective': {'target_role': value},
    }
    if section in field_map:
        for field, val in field_map[section].items():
            setattr(profile, field, val)
        profile.save()
        return JsonResponse({'success': True})

    if section == 'skills':
        from apps.assessment.models import UserSkill
        skill_name = data.get('skill_name', '').strip()
        proficiency = data.get('proficiency', 3)
        category = data.get('category', 'technical')
        if skill_name:
            obj, created = UserSkill.objects.get_or_create(
                user=request.user, skill_name=skill_name,
                defaults={'proficiency': int(proficiency), 'category': category}
            )
            if not created:
                obj.proficiency = int(proficiency)
                obj.category = category
                obj.save()
            return JsonResponse({'success': True, 'created': created})
        return JsonResponse({'error': 'Skill name required'}, status=400)

    if section == 'experience':
        profile.current_role = data.get('title', profile.current_role)
        profile.years_experience = data.get('years', profile.years_experience)
        profile.save()
        return JsonResponse({'success': True})

    if section == 'education':
        profile.education_level = data.get('level', profile.education_level)
        profile.field_of_study = data.get('field', profile.field_of_study)
        profile.institution = data.get('institution', profile.institution)
        profile.graduation_year = data.get('year', profile.graduation_year)
        profile.save()
        return JsonResponse({'success': True})

    if section == 'certification':
        from apps.assessment.models import UserSkill
        name = data.get('name', '').strip()
        if name:
            UserSkill.objects.get_or_create(
                user=request.user, skill_name=name,
                defaults={'category': 'certification', 'proficiency': 3}
            )
            return JsonResponse({'success': True})
        return JsonResponse({'error': 'Name required'}, status=400)

    return JsonResponse({'error': 'Unknown section'}, status=400)


@login_required
@xframe_options_exempt
def cv_render_view(request, pk):
    """Render original CV in browser (PDF embedded, DOCX→HTML)."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    file_path = cv_upload.file.path

    if not os.path.exists(file_path):
        return HttpResponse(
            '<div style="display:flex;align-items:center;justify-content:center;'
            'height:100vh;font-family:sans-serif;color:#666;background:#f8f9fa;">'
            '<div style="text-align:center;"><h2 style="color:#dc3545;">File Not Found</h2>'
            f'<p>The uploaded file "{cv_upload.original_filename}" is missing from storage.</p>'
            '<p class="small">Please re-upload the document.</p></div></div>'
        )

    try:
        if cv_upload.file_type == 'pdf':
            with open(file_path, 'rb') as f:
                response = HttpResponse(f.read(), content_type='application/pdf')
            response['Content-Disposition'] = 'inline; filename="' + cv_upload.original_filename + '"'
            return response
    except Exception as e:
        return HttpResponse(
            '<div style="display:flex;align-items:center;justify-content:center;'
            'height:100vh;font-family:sans-serif;color:#666;background:#f8f9fa;">'
            '<div style="text-align:center;"><h2 style="color:#dc3545;">Error Reading PDF</h2>'
            f'<p>Could not read the PDF file: {str(e)}</p>'
            '<p class="small">Try re-uploading the document.</p></div></div>'
        )

    # DOCX → HTML via python-docx preserving formatting
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    try:
        doc = Document(file_path)
    except Exception as e:
        return HttpResponse(
            '<div style="display:flex;align-items:center;justify-content:center;'
            'height:100vh;font-family:sans-serif;color:#666;background:#f8f9fa;">'
            '<div style="text-align:center;"><h2 style="color:#dc3545;">Error Reading DOCX</h2>'
            f'<p>Could not parse the DOCX file: {str(e)}</p>'
            '<p class="small">The file may be corrupted. Try re-uploading.</p></div></div>'
        )
    parts = []

    # Paragraph styles mapping
    style_map = {
        'Title': ('h1', ''),
        'Subtitle': ('h2', 'color:#6c757d;'),
        'Heading': ('h2', ''),
        'Heading 1': ('h1', ''),
        'Heading 2': ('h2', ''),
        'Heading 3': ('h3', ''),
        'Heading 4': ('h4', ''),
        'List Bullet': ('p', 'padding-left:20px;text-indent:-10px;'),
        'List Number': ('p', 'padding-left:20px;'),
        'Normal': ('p', ''),
    }

    def render_run(run):
        text = run.text
        if not text:
            return ''
        tag = 'span'
        style = ''
        if run.bold:
            tag = 'strong'
        if run.italic:
            tag = 'em' if tag == 'span' else tag
            style += 'font-style:italic;'
        if run.underline:
            style += 'text-decoration:underline;'
        if style:
            return f'<{tag} style="{style}">{text}</{tag}>'
        return f'<{tag}>{text}</{tag}>'

    in_list = False

    def render_para(para):
        nonlocal in_list
        style_name = para.style.name if para.style else 'Normal'
        tag, extra_style = style_map.get(style_name, ('p', ''))

        if 'List' in style_name:
            bullet_char = '•' if 'Bullet' in style_name else '1.'
            runs_html = ''.join(render_run(r) for r in para.runs)
            text = runs_html or para.text
            if not text.strip():
                return '<br>'
            if not in_list:
                in_list = True
                return '<ul style="list-style:none;padding-left:20px;margin:2px 0;">' \
                       f'<li style="margin:1px 0;">{bullet_char} {text}</li>'
            return f'<li style="margin:1px 0;">{bullet_char} {text}</li>'

        if in_list:
            in_list = False
            parts.append('</ul>')

        runs_html = ''.join(render_run(r) for r in para.runs)
        text = runs_html or para.text
        if not text.strip():
            return '<br>'

        align = para.alignment
        align_style = ''
        if align == WD_ALIGN_PARAGRAPH.CENTER:
            align_style = 'text-align:center;'
        elif align == WD_ALIGN_PARAGRAPH.RIGHT:
            align_style = 'text-align:right;'

        if tag in ('h1', 'h2', 'h3', 'h4'):
            return f'<{tag} style="{extra_style}{align_style}margin:8px 0 4px;">{text}</{tag}>'
        return f'<{tag} style="{extra_style}{align_style}margin:3px 0;">{text}</{tag}>'

    for para in doc.paragraphs:
        parts.append(render_para(para))
    if in_list:
        parts.append('</ul>')

    # Tables
    for table in doc.tables:
        parts.append(
            '<div style="overflow-x:auto;margin:12px 0;">'
            '<table style="border-collapse:collapse;width:100%;font-size:0.88rem;">'
        )
        for row_idx, row in enumerate(table.rows):
            tag = 'th' if row_idx == 0 else 'td'
            bg = 'background:#f0f0f0;' if row_idx == 0 else ''
            parts.append('<tr>')
            for cell in row.cells:
                parts.append(
                    f'<{tag} style="border:1px solid #ccc;padding:6px 10px;{bg}'
                    f'text-align:left;vertical-align:top;">{cell.text}</{tag}>'
                )
            parts.append('</tr>')
        parts.append('</table></div>')

    content = '\n'.join(parts)
    html = f'''<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
*{{margin:0;padding:0;box-sizing:border-box;}}
body{{font-family:'Calibri','Segoe UI',Arial,sans-serif;
padding:48px 56px;max-width:700px;margin:0 auto;
color:#000;line-height:1.45;font-size:10.5pt;}}
h1{{font-size:16pt;font-weight:700;margin:14px 0 4px;color:#1a1a2e;letter-spacing:-0.01em;}}
h2{{font-size:13pt;font-weight:600;margin:12px 0 3px;color:#1a1a2e;}}
h3{{font-size:11.5pt;font-weight:600;margin:10px 0 3px;color:#1a1a2e;}}
h4{{font-size:10.5pt;font-weight:600;margin:8px 0 2px;color:#1a1a2e;}}
p{{margin:2px 0;font-size:10.5pt;line-height:1.45;}}
br{{display:block;margin:3px 0;}}
ul{{list-style:none;padding-left:18px;margin:2px 0;}}
li{{margin:1px 0;font-size:10.5pt;line-height:1.45;}}
hr{{border:none;border-top:1px solid #ddd;margin:12px 0;}}
table{{margin:10px 0;border-collapse:collapse;width:100%;}}
th{{font-weight:600;background:#f5f5f5;font-size:9.5pt;}}
td,th{{border:1px solid #ccc;padding:5px 8px;text-align:left;vertical-align:top;font-size:9.5pt;}}
strong{{font-weight:700;}} em{{font-style:italic;}}
</style></head>
<body>{content}</body></html>'''
    return HttpResponse(html)


@login_required
def cv_history_view(request):
    """List all uploaded CVs."""
    cvs = CVUpload.objects.filter(user=request.user, is_deleted=False).select_related('analysis').order_by('-uploaded_at')
    context = {
        'page_title': 'CV History',
        'cvs': cvs,
    }
    return render(request, 'cv/history.html', context)


@login_required
def cv_reanalyze_view(request, pk):
    """Re-run analysis on an existing CV."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)

    if not cv_upload.parsed_text:
        try:
            cv_upload.parsed_text = parse_cv_file(cv_upload.file.path)
            cv_upload.save()
        except Exception as e:
            messages.error(request, f'Failed to re-parse CV: {str(e)}')
            return redirect('cv_analyzer:history')

    analysis_result = analyze_cv(cv_upload.parsed_text)

    # Update or create analysis
    CVAnalysis.objects.update_or_create(
        cv_upload=cv_upload,
        defaults={
            'ats_score': analysis_result['ats_score'],
            'grammar_score': analysis_result['grammar_score'],
            'keyword_score': analysis_result['keyword_score'],
            'format_score': analysis_result['format_score'],
            'overall_score': analysis_result['overall_score'],
            'extracted_sections': analysis_result['extracted_sections'],
            'extracted_skills': analysis_result['extracted_skills'],
            'detected_keywords': analysis_result.get('detected_keywords', analysis_result.get('extracted_skills', [])),
            'contact_info': analysis_result.get('contact_info', {}),
            'suggestions': analysis_result['suggestions'],
            'missing_sections': analysis_result['missing_sections'],
            'improvements': analysis_result['improvements'],
            'ai_feedback': analysis_result['ai_feedback'],
        }
    )

    try:
        from apps.admin_panel.utils import log_activity
        log_activity(request.user, 'update', f'Re-analyzed CV "{cv_upload.original_filename}"', 'CVAnalysis', cv_upload.pk, request)
    except Exception:
        pass

    messages.success(request, 'CV re-analyzed successfully!')
    return redirect('cv_analyzer:analysis', pk=cv_upload.pk)


@login_required
def export_report_view(request, pk):
    """Export CV analysis report as PDF."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)

    if not cv_upload.has_analysis:
        messages.error(request, 'No analysis found for this CV.')
        return redirect('cv_analyzer:history')

    analysis = cv_upload.analysis

    try:
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import inch
        import io

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72)
        styles = getSampleStyleSheet()
        story = []

        # Custom styles
        title_style = ParagraphStyle('CustomTitle', parent=styles['Title'], fontSize=20, textColor=colors.HexColor('#1a8f5c'))
        heading_style = ParagraphStyle('Heading', parent=styles['Heading2'], textColor=colors.HexColor('#1a8f5c'))

        # Title
        story.append(Paragraph('CV Analysis Report', title_style))
        story.append(Paragraph(f'File: {cv_upload.original_filename}', styles['Normal']))
        story.append(Paragraph(f'Generated: {analysis.analyzed_at.strftime("%B %d, %Y")}', styles['Normal']))
        story.append(Spacer(1, 0.3 * inch))

        # Scores table
        story.append(Paragraph('Analysis Scores', heading_style))
        score_data = [
            ['Metric', 'Score', 'Rating'],
            ['Overall Score', f'{analysis.overall_score}/100', _get_score_label(analysis.overall_score)[0]],
            ['ATS Compatibility', f'{analysis.ats_score}/100', _get_score_label(analysis.ats_score)[0]],
            ['Grammar & Professionalism', f'{analysis.grammar_score}/100', _get_score_label(analysis.grammar_score)[0]],
            ['Keyword Optimization', f'{analysis.keyword_score}/100', _get_score_label(analysis.keyword_score)[0]],
            ['Format & Structure', f'{analysis.format_score}/100', _get_score_label(analysis.format_score)[0]],
        ]
        table = Table(score_data, colWidths=[2.5 * inch, 1.5 * inch, 1.5 * inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a8f5c')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f8f4')]),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.3 * inch))

        # Skills
        if analysis.extracted_skills:
            story.append(Paragraph('Detected Skills', heading_style))
            skills_text = ' • '.join(analysis.extracted_skills)
            story.append(Paragraph(skills_text, styles['Normal']))
            story.append(Spacer(1, 0.2 * inch))

        # Suggestions
        if analysis.suggestions:
            story.append(Paragraph('Improvement Suggestions', heading_style))
            for i, sug in enumerate(analysis.suggestions, 1):
                if isinstance(sug, dict):
                    story.append(Paragraph(f'<b>{i}. [{sug.get("category", "General")}]</b> {sug.get("suggestion", "")}', styles['Normal']))
                    story.append(Spacer(1, 0.1 * inch))

        doc.build(story)
        buffer.seek(0)

        response = HttpResponse(buffer, content_type='application/pdf')
        filename = f'cv_analysis_{cv_upload.pk}.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except ImportError:
        messages.error(request, 'PDF export requires ReportLab. Please install it: pip install reportlab')
        return redirect('cv_analyzer:analysis', pk=pk)


@login_required
def cv_delete_view(request, pk):
    """Delete a CV upload."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    if request.method == 'POST':
        try:
            from apps.admin_panel.utils import log_activity
            log_activity(request.user, 'delete', f'Deleted CV "{cv_upload.original_filename}"', 'CVUpload', cv_upload.pk, request)
        except Exception:
            pass
        cv_upload.delete()
        messages.success(request, 'CV deleted successfully.')
    
    referer = request.META.get('HTTP_REFERER')
    if referer:
        return redirect(referer)
    return redirect('cv_analyzer:history')


@login_required
def set_active_cv_view(request, pk):
    """Set a CV as the active/primary CV."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    CVUpload.objects.filter(user=request.user).update(is_active=False)
    cv_upload.is_active = True
    cv_upload.save()
    messages.success(request, f'{cv_upload.original_filename} is now your active CV.')
    return redirect('cv_analyzer:history')


def _get_score_label(score):
    """Return (label, bootstrap color) for a score."""
    if score >= 80:
        return ('Excellent', 'success')
    elif score >= 60:
        return ('Good', 'primary')
    elif score >= 40:
        return ('Fair', 'warning')
    else:
        return ('Poor', 'danger')


@login_required
def cv_download_view(request, pk):
    """Download the original uploaded CV file."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    import os
    if os.path.exists(cv_upload.file.path):
        with open(cv_upload.file.path, 'rb') as f:
            response = HttpResponse(f.read(), content_type='application/octet-stream')
            response['Content-Disposition'] = f'attachment; filename="{cv_upload.original_filename}"'
            return response
    messages.error(request, 'File not found on disk.')
    return redirect('cv_analyzer:history')


@login_required
def cv_compare_view(request):
    """Compare multiple CV versions."""
    cvs = CVUpload.objects.filter(user=request.user, is_deleted=False).select_related('analysis')
    cv_ids = request.GET.getlist('cv')
    selected_cvs = []
    if cv_ids:
        selected_cvs = CVUpload.objects.filter(pk__in=cv_ids, user=request.user, is_deleted=False).select_related('analysis')
    context = {
        'page_title': 'Compare CV Versions',
        'cvs': cvs,
        'selected_cvs': selected_cvs,
    }
    return render(request, 'cv/compare.html', context)


def _collect_resume_data(user, cv_upload, analysis):
    """Collect user profile + CV data for resume template rendering."""
    from apps.assessment.models import UserSkill
    profile = getattr(user, 'profile', None)
    skills_qs = UserSkill.objects.filter(user=user)
    all_skills = list(skills_qs.values_list('skill_name', flat=True))
    extracted = analysis.extracted_skills or []
    if isinstance(extracted, str):
        extracted = [s.strip() for s in extracted.split(',') if s.strip()]
    all_skills += [s.strip() for s in extracted if s.strip() and s.strip() not in all_skills]
    contact = analysis.contact_info or {}
    sections = analysis.extracted_sections or {}
    experience_items = []
    raw_exp = sections.get('experience')
    if raw_exp:
        if isinstance(raw_exp, str):
            raw_exp = [{'title': raw_exp.strip(), 'company': '', 'dates': '', 'description': ''}]
        for exp in raw_exp:
            if isinstance(exp, dict):
                experience_items.append(exp)
            elif isinstance(exp, str):
                experience_items.append({'title': exp, 'company': '', 'dates': '', 'description': ''})
    if not experience_items and profile and profile.current_role:
        experience_items.append({
            'title': profile.current_role,
            'company': '',
            'dates': '',
            'description': profile.bio or '',
        })
    certifications = []
    raw_cert = sections.get('certifications', '')
    if raw_cert:
        if isinstance(raw_cert, str):
            certifications = [c.strip() for c in re.split(r'[,;]', raw_cert) if c.strip()]
        elif isinstance(raw_cert, list):
            certifications = [str(c) for c in raw_cert if str(c).strip()]
    extra_sections = []
    displayed_keys = {'contact', 'summary', 'education', 'experience', 'skills', 'certifications'}
    for key, val in sections.items():
        if key not in displayed_keys and val and (isinstance(val, str) and val.strip()):
            extra_sections.append({'name': key.replace('_', ' ').title(), 'content': val})
    return {
        'user': user,
        'profile': profile,
        'cv': cv_upload,
        'analysis': analysis,
        'all_skills': all_skills,
        'certifications': certifications,
        'contact': contact,
        'sections': sections,
        'experience_items': experience_items,
        'extra_sections': extra_sections,
    }


@login_required
def resume_templates_view(request):
    """List available resume templates and preview/select."""
    cvs = CVUpload.objects.filter(user=request.user, is_deleted=False).select_related('analysis').order_by('-uploaded_at')
    selected_cv_pk = request.GET.get('cv')
    template = request.GET.get('template', 'modern')
    selected_cv = None
    if selected_cv_pk:
        selected_cv = get_object_or_404(CVUpload, pk=selected_cv_pk, user=request.user)
    templates_info = [
        ('modern', 'Modern', 'Clean, minimal design with a blue header. Perfect for tech roles.', 'fas fa-leaf'),
        ('professional', 'Professional', 'Classic serif layout for traditional industries.', 'fas fa-briefcase'),
        ('creative', 'Creative', 'Bold sidebar design with accent colors. Great for creative roles.', 'fas fa-paint-brush'),
        ('elegant', 'Elegant', 'Two-column beautiful modern design tailored for professional layouts.', 'fas fa-star text-warning'),
    ]
    context = {
        'page_title': 'Resume Templates',
        'cvs': cvs,
        'selected_cv': selected_cv,
        'template': template,
        'templates_info': templates_info,
    }
    return render(request, 'cv/resume_templates.html', context)


@login_required
@xframe_options_exempt
@xframe_options_exempt
def resume_preview_view(request, pk, template_name):
    """Preview CV data rendered in a selected template."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    analysis = getattr(cv_upload, 'analysis', None)
    if not analysis:
        messages.error(request, 'This CV has not been analyzed yet.')
        return redirect('cv_analyzer:resume_templates')
    data = _collect_resume_data(request.user, cv_upload, analysis)
    return render(request, f'resume_templates/{template_name}.html', data)


@login_required
@xframe_options_exempt
def resume_download_pdf_view(request, pk, template_name):
    """Generate PDF from CV data using xhtml2pdf."""
    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    analysis = getattr(cv_upload, 'analysis', None)
    if not analysis:
        messages.error(request, 'This CV has not been analyzed yet.')
        return redirect('cv_analyzer:resume_templates')

    data = _collect_resume_data(request.user, cv_upload, analysis)
    data['is_pdf'] = True
    from django.template.loader import render_to_string
    html_string = render_to_string(f'resume_templates/{template_name}.html', data)

    try:
        from xhtml2pdf import pisa
        response = HttpResponse(content_type='application/pdf')
        filename = f"Resume_{request.user.last_name or request.user.first_name}_{template_name}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        pisa_status = pisa.CreatePDF(html_string, dest=response, link_callback=fetch_resources)
        if pisa_status.err:
            raise Exception('PDF generation failed')
        return response
    except Exception as e:
        messages.error(request, f'Could not generate PDF: {e}. Using browser print instead.')
        return render(request, f'resume_templates/{template_name}.html', data)


def fetch_resources(uri, rel):
    """Callback for xhtml2pdf to resolve static files."""
    import os
    from django.conf import settings
    if uri.startswith(settings.STATIC_URL):
        path = os.path.join(settings.STATIC_ROOT or settings.BASE_DIR, uri.replace(settings.STATIC_URL, ''))
        if os.path.exists(path):
            return path
    if uri.startswith(settings.MEDIA_URL):
        path = os.path.join(settings.MEDIA_ROOT, uri.replace(settings.MEDIA_URL, ''))
        if os.path.exists(path):
            return path
    return uri


def _create_cv_notification(user, cv_upload, score):
    """Create a notification after CV analysis."""
    try:
        from apps.notifications.models import Notification
        label = _get_score_label(score)[0]
        Notification.objects.create(
            user=user,
            title='CV Analysis Complete',
            message=f'Your CV "{cv_upload.original_filename}" has been analyzed. Overall score: {score}/100 ({label}).',
            notification_type='cv',
            link=f'/cv/{cv_upload.pk}/analysis/',
        )
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Feature 2: One-Click CV Bullet Optimizer
# ---------------------------------------------------------------------------

@login_required
def optimize_bullets_view(request, pk, job_match_pk):
    """
    Show and process AI-powered bullet point optimization for a CV
    against the missing keywords from a specific job match.
    """
    from .utils.ai_service import optimize_cv_bullets
    from apps.job_matching.models import JobMatch

    cv_upload = get_object_or_404(CVUpload, pk=pk, user=request.user)
    job_match = get_object_or_404(JobMatch, pk=job_match_pk, user=request.user)

    optimization = None
    if request.method == 'POST':
        missing_keywords = job_match.missing_keywords or []
        job_title = job_match.job_description.title
        optimization = optimize_cv_bullets(cv_upload.parsed_text, missing_keywords, job_title)

    context = {
        'page_title': 'Optimize CV Bullets',
        'cv_upload': cv_upload,
        'job_match': job_match,
        'optimization': optimization,
    }
    return render(request, 'cv/optimize_bullets.html', context)


# Force django auto-reload to pick up xhtml2pdf 0.2.21
