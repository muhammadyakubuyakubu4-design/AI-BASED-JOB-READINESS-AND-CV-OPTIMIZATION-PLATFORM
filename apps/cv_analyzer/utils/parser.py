"""
CV Parser - Extracts text and sections from PDF/DOCX files.
"""
import re
import os


def parse_cv_file(file_path: str) -> str:
    """Parse a CV file (PDF or DOCX) and return its text content."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return _parse_pdf(file_path)
    elif ext == '.docx':
        return _parse_docx(file_path)
    else:
        raise ValueError(f'Unsupported file type: {ext}')


def _parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    try:
        import PyPDF2
        text = ''
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + '\n'
        return text.strip()
    except Exception as e:
        raise RuntimeError(f'Failed to parse PDF: {str(e)}')


def _parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return '\n'.join(paragraphs)
    except Exception as e:
        raise RuntimeError(f'Failed to parse DOCX: {str(e)}')


def extract_sections(text: str) -> dict:
    """
    Extract major CV sections from parsed text.
    Returns a dictionary with section names as keys and content as values.
    """
    sections = {
        'contact': '',
        'summary': '',
        'education': '',
        'experience': '',
        'skills': '',
        'certifications': '',
        'projects': '',
        'languages': '',
        'interests': '',
    }

    # Section header patterns
    section_patterns = {
        'contact': r'(?i)^\s*(contact|personal\s+info(rmation)?|personal\s+details|contact\s+details|contact\s+information|profile\s+contact)\s*[:.]?\s*$',
        'summary': r'(?i)^\s*(summary|objective|career\s+objective|professional\s+summary|professional\s+profile|profile|about\s+me|accomplishment\s+summary|personal\s+statement|qualifications?\s+summary)\s*[:.]?\s*$',
        'education': r'(?i)^\s*(education|academic\s+background|academic\s+qualifications?|educational\s+qualifications?|education\s+and\s+training|qualifications?|degrees?|universit(y|ies)\s+attended)\s*[:.]?\s*$',
        'experience': r'(?i)^\s*(experience|professional\s+experience|work\s+experience|employment\s+history|employment|career\s+history|work\s+history|work\s+experience\s+history|relevant\s+experience|professional\s+background|experience\s+history)\s*[:.]?\s*$',
        'skills': r'(?i)^\s*(skills?|technical\s+skills?(\s+(&|and)\s+tools?)?|core\s+competencies?|competencies?|abilities?|professional\s+skills?|key\s+skills?|skills\s+(and|&)\s+competencies?|areas?\s+of\s+expertise|expertise|software\s+skills?|computer\s+skills?|hard\s+skills?|soft\s+skills?|tools?\s+(and|&)\s+technologies?|skills?\s+(and|&)\s+abilities?|skills?\s+(and|&)\s+knowledge|technical\s+(skills|knowledge))(\s+(&|and)\s+competencies?)?\s*[:.]?\s*$',
        'certifications': r'(?i)^\s*(certifications?|certificates?|certification\s*(&|and)\s*training|professional\s+certifications?|certifications?\s*(&|and)\s*(licenses?|accreditations?|training|courses?)|accreditations?|licens(es|ures?)|licenses?\s*(&|and)\s*certifications?|professional\s+development)\s*[:.]?\s*$',
        'projects': r'(?i)^\s*(projects?|personal\s+projects?|major\s+projects?|academic\s+projects?|key\s+projects?|project\s+work|portfolio|selected\s+projects?|sample\s+projects?|major\s+projects?\s*(&|and)\s*(portfolio|others?|projects?)|projects?\s*(&|and)\s*(portfolio|achievements?|case\s+studies?)|projects\s+undertaken|notable\s+projects?)\s*[:.]?\s*$',
        'languages': r'(?i)^\s*(languages?|languages?\s+spoken|spoken\s+languages?|linguistic\s+abilities?|language\s+proficiency|foreign\s+languages?)\s*[:.]?\s*$',
        'interests': r'(?i)^\s*(interests?|hobbies?|activities?|interest\s+and\s+hobbies?|extracurricular\s+(activities?|interests?)|additional\s+interests?|personal\s+interests?)\s*[:.]?\s*$',
    }

    lines = text.split('\n')
    current_section = 'contact'
    section_content = {k: [] for k in sections}

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # Check if this line is a section header
        matched_section = None
        for section_name, pattern in section_patterns.items():
            if re.match(pattern, line) and len(line) < 60:
                matched_section = section_name
                break

        if matched_section:
            current_section = matched_section
        else:
            section_content[current_section].append(line)

    # Join section content
    for section_name in sections:
        sections[section_name] = '\n'.join(section_content[section_name])

    return sections


def extract_contact_info(text: str) -> dict:
    """Extract contact information from CV text."""
    contact = {
        'email': '',
        'phone': '',
        'linkedin': '',
        'github': '',
        'website': '',
        'location': '',
    }

    # Email pattern
    email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    if email_match:
        contact['email'] = email_match.group()

    # Phone pattern (various formats)
    phone_match = re.search(
        r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text
    )
    if phone_match:
        contact['phone'] = phone_match.group()

    # LinkedIn
    linkedin_match = re.search(r'linkedin\.com/in/[\w-]+', text, re.IGNORECASE)
    if linkedin_match:
        contact['linkedin'] = linkedin_match.group()

    # GitHub
    github_match = re.search(r'github\.com/[\w-]+', text, re.IGNORECASE)
    if github_match:
        contact['github'] = github_match.group()

    return contact


def extract_skills_from_text(text: str) -> list:
    """Extract skill keywords from CV text using a predefined skill database."""
    common_skills = [
        # Programming Languages
        'Python', 'Java', 'JavaScript', 'TypeScript', 'C++', 'C#', 'C', 'Ruby',
        'PHP', 'Swift', 'Kotlin', 'Go', 'Rust', 'Scala', 'R', 'MATLAB',
        # Web Technologies
        'HTML', 'CSS', 'React', 'Angular', 'Vue', 'Node.js', 'Express', 'Django',
        'Flask', 'FastAPI', 'Bootstrap', 'jQuery', 'REST', 'GraphQL', 'WebSocket',
        # Databases
        'MySQL', 'PostgreSQL', 'MongoDB', 'SQLite', 'Redis', 'Elasticsearch',
        'Oracle', 'SQL Server', 'DynamoDB', 'Firebase',
        # Cloud & DevOps
        'AWS', 'Azure', 'GCP', 'Docker', 'Kubernetes', 'CI/CD', 'Jenkins',
        'Terraform', 'Ansible', 'Linux', 'Git', 'GitHub', 'GitLab',
        # Data Science & AI
        'Machine Learning', 'Deep Learning', 'NLP', 'TensorFlow', 'PyTorch',
        'Scikit-learn', 'Pandas', 'NumPy', 'Matplotlib', 'Data Analysis',
        'Power BI', 'Tableau', 'Excel', 'Statistics',
        # Finance & Accounting
        'Accounting', 'Auditing', 'Bookkeeping', 'Financial Reporting', 'Taxation',
        'Financial Analysis', 'Payroll', 'Reconciliation', 'QuickBooks', 'SAP',
        'ERP', 'Accounts Payable', 'Accounts Receivable', 'Budgeting', 'Forecasting',
        'Compliance', 'Risk Management', 'IFRS', 'GAAP', 'Financial Statements',
        # Soft Skills
        'Leadership', 'Communication', 'Teamwork', 'Problem Solving',
        'Project Management', 'Agile', 'Scrum', 'Time Management',
    ]

    found_skills = []
    text_upper = text.upper()

    for skill in common_skills:
        if skill.upper() in text_upper:
            found_skills.append(skill)

    return list(set(found_skills))
